#!/usr/bin/env python3
"""Lint a stagecraft deck source and assemble the single-file deck.

usage: python build.py deck.src.html [-o deck.html] [--brief brief.json] [--offline] [--lint-only] [--force]

Errors block the build (exit 1). Warnings print but do not block; fix them unless there is a reason not to.
"""
import argparse
import json
import re
import sys
import urllib.request
from html.parser import HTMLParser
from pathlib import Path

SKILL = Path(__file__).resolve().parent.parent
GSAP_VER = "3.13.0"
GSAP_FILES = ["gsap.min.js", "SplitText.min.js", "ScrambleTextPlugin.min.js"]
GSAP_CDN = f"https://cdn.jsdelivr.net/npm/gsap@{GSAP_VER}/dist/"

VOID = {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "source", "track", "wbr"}
CODE_TAGS = {"code", "pre", "kbd", "samp"}
TEXT_ATTRS = ("data-title", "data-tip", "alt", "aria-label", "title")
BLOCK = {"p", "li", "h1", "h2", "h3", "h4", "h5", "h6", "div", "td", "th", "figcaption", "blockquote", "button", "dt", "dd", "text", "tspan", "label"}

DASHES = re.compile(r"[‒–—―]")
SPACED_HYPHEN = re.compile(r"\w\s+-{1,2}\s+\w")
EMOJI = re.compile("[\U0001F300-\U0001FAFF\U00002600-\U000027BF\U0001F000-\U0001F2FF️]")
NUMBER = re.compile(r"(?<![A-Za-z0-9\-_/#.])\d+(?:[.,]\d+)*")
ANY_NUMBER = re.compile(r"\d+(?:[.,]\d+)*")
BANNED_FONTS = ["inter", "roboto", "arial", "helvetica", "open sans", "lato", "montserrat", "poppins", "space grotesk", "segoe ui", "system-ui"]
GENERIC_TITLES = {"introduction", "intro", "overview", "agenda", "summary", "conclusion", "conclusions", "thank you", "thanks",
                  "q&a", "questions", "key takeaways", "takeaways", "background", "next steps", "about us", "the problem", "the solution"}
BANNED_PHRASES = [
    "delve", "leverage", "leveraging", "seamless", "seamlessly", "unlock", "unlocking", "elevate", "game-changer", "game changer",
    "revolutionize", "revolutionary", "tapestry", "realm", "harness the power", "cutting-edge", "cutting edge", "fast-paced",
    "ever-evolving", "ever-changing", "landscape", "empower", "supercharge", "next-level", "embark", "testament", "paradigm shift",
    "synergy", "robust solution", "in today's", "dive into", "deep dive", "let's explore", "at the end of the day", "it's worth noting",
    "furthermore", "moreover", "in conclusion", "a myriad of", "plethora", "transformative", "unleash", "journey", "holistic",
    "best-in-class", "world-class", "state-of-the-art", "mission-critical", "pivotal", "crucial role", "navigate the",
]
NOT_X_BUT_Y = re.compile(r"\b(not just|not only|isn't just|is not just|more than just|it's not about)\b[^.]{0,80}\b(but|it's|it is)\b", re.I)


class DeckParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.stack = []          # (tag, attrs)
        self.slides = []         # dicts
        self.cur = None
        self.fonts_css = []
        self.in_style = False
        self.allow_fonts = set()
        self.has_stage = self.has_viewport = False

    def _skip(self):
        for tag, a in self.stack:
            if tag in ("script", "style", "template") or "data-deco" in a or a.get("aria-hidden") == "true":
                return True
        return False

    def _in(self, names):
        return any(t in names for t, _ in self.stack)

    def _in_notes(self):
        return any("notes" in (a.get("class") or "").split() for _, a in self.stack)

    def handle_starttag(self, tag, attrs):
        a = dict((k, v if v is not None else "") for k, v in attrs)
        cls = (a.get("class") or "").split()
        if "kx-stage" in cls: self.has_stage = True
        if "kx-viewport" in cls: self.has_viewport = True
        if tag == "meta" and a.get("name") == "kx-allow-fonts":
            self.allow_fonts |= {f.strip().lower() for f in a.get("content", "").split(",") if f.strip()}
        if tag == "link" and "fonts.googleapis.com" in a.get("href", ""):
            self.fonts_css.extend(re.findall(r"family=([^&:]+)", a["href"].replace("+", " ")))
        if tag == "section" and "slide" in cls:
            self.cur = {"id": a.get("id"), "layout": a.get("data-layout"), "claims": (a.get("data-claims") or "").split(),
                        "title_attr": a.get("data-title"), "text": [], "notes": [], "code": [], "headings": [],
                        "anims": 0, "has_notes": False, "line": self.getpos()[0]}
            self.slides.append(self.cur)
        if self.cur is not None:
            if "data-anim" in a: self.cur["anims"] += 1
            if "notes" in cls: self.cur["has_notes"] = True
            for k in TEXT_ATTRS:
                if a.get(k) and not self._skip():
                    self.cur["text"].append("\n" + a[k] + "\n")
        if tag == "style": self.in_style = True
        if self.cur is not None and tag in BLOCK:
            self.cur["text"].append("\n"); self.cur["notes"].append("\n")
        if tag not in VOID:
            self.stack.append((tag, a))
        if tag in ("h1", "h2") and self.cur is not None:
            self.cur["headings"].append("")

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if tag not in VOID and self.stack and self.stack[-1][0] == tag:
            self.stack.pop()

    def handle_endtag(self, tag):
        if tag == "style": self.in_style = False
        if self.cur is not None and tag in BLOCK:
            self.cur["text"].append("\n"); self.cur["notes"].append("\n")
        for i in range(len(self.stack) - 1, -1, -1):
            if self.stack[i][0] == tag:
                del self.stack[i:]
                break
        if tag == "section" and self.cur is not None and not any(t == "section" for t, _ in self.stack):
            self.cur = None

    def handle_data(self, data):
        if self.in_style:
            self.fonts_css.extend(re.findall(r"font-family\s*:\s*([^;}{]+)", data))
            self.fonts_css.extend(re.findall(r"--font[\w-]*\s*:\s*([^;}{]+)", data))
            return
        if self.cur is None or not data.strip() or self._skip():
            return
        if self._in_notes():
            self.cur["notes"].append(data)
        elif self._in(CODE_TAGS):
            self.cur["code"].append(data)
        else:
            self.cur["text"].append(data)
            if self._in(("h1", "h2")) and self.cur["headings"]:
                self.cur["headings"][-1] += data


def source_text(path, cache):
    """Whitespace-normalized text of a local source, or None if it cannot be read."""
    if path in cache:
        return cache[path]
    text = None
    try:
        if not path.is_file():
            text = None
        elif path.suffix.lower() == ".docx":
            import docx
            d = docx.Document(str(path))
            parts = [p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells]
            text = "\n".join(parts)
        elif path.suffix.lower() == ".pdf":
            import pypdf
            text = "\n".join(pg.extract_text() or "" for pg in pypdf.PdfReader(str(path)).pages)
        elif path.suffix.lower() in (".pptx",):
            from pptx import Presentation
            text = "\n".join(sh.text_frame.text for sl in Presentation(str(path)).slides for sh in sl.shapes if sh.has_text_frame)
        elif path.suffix.lower() not in (".png", ".jpg", ".jpeg", ".gif", ".webp", ".mp4"):
            text = path.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        text = None
    cache[path] = ws(text) if text is not None else None
    return cache[path]


def norm_num(s):
    return s.replace(",", "")


def ws(s):
    return re.sub(r"\s+", " ", s).strip()


def lint(src_html, brief, brief_dir):
    errors, warnings, verify = [], [], []
    p = DeckParser()
    p.feed(src_html)

    if not (p.has_viewport and p.has_stage):
        errors.append("structure: deck needs <div class=\"kx-viewport\"><div class=\"kx-stage\"> wrapping the slides")
    if not p.slides:
        errors.append("structure: no <section class=\"slide\"> found")

    # fonts
    fam = " | ".join(p.fonts_css).lower()
    for f in BANNED_FONTS:
        if f in p.allow_fonts:
            continue
        if re.search(r"(^|[\s,'\"|])" + re.escape(f) + r"($|[\s,'\"|;])", fam):
            first_choice = any(re.match(r"\s*['\"]?" + re.escape(f) + r"\b", x.lower()) for x in p.fonts_css)
            if first_choice or f in ("inter", "space grotesk", "poppins", "montserrat"):
                errors.append(f"font: '{f}' is a default-AI tell. Pick a distinctive face (references/design.md) or allow it via <meta name=\"kx-allow-fonts\" content=\"{f}\"> when a brand book requires it")

    # ledger
    claims, srcs, policy, cache, unverified = {}, {}, "strict", {}, set()
    if brief:
        policy = brief.get("source_policy", "strict")
        srcs = {s["id"]: s for s in brief.get("sources", [])}
        claims = {c["id"]: c for c in brief.get("claims", [])}
        for c in claims.values():
            s = srcs.get(c.get("source"))
            if not s:
                errors.append(f"ledger: claim {c['id']} cites unknown source '{c.get('source')}'")
                continue
            if s.get("kind") == "model":
                if policy == "strict":
                    errors.append(f"ledger: claim {c['id']} uses model knowledge but source_policy is strict")
                else:
                    verify.append(f"{c['id']}: {c.get('text')}")
            ref, q = s.get("ref", ""), c.get("quote")
            if not q or s.get("kind") in ("user", "model") or re.match(r"^[a-z]+://", ref or ""):
                continue
            body = source_text(brief_dir / ref, cache)
            if body is None:
                unverified.add(f"{s['id']} ({ref})")
            elif ws(q) not in body:
                errors.append(f"ledger: claim {c['id']} quote not found verbatim in {ref}")
        for u in sorted(unverified):
            warnings.append(f"ledger: quotes from source {u} were not verified (file missing or unreadable format). Use a path relative to brief.json or absolute, or extract the text to a .txt source")
    else:
        warnings.append("ledger: no brief.json found, so slide content is not checked against sources")
    ledger_text = " ".join(norm_num(c.get("text", "") + " " + (c.get("quote") or "")) for c in claims.values())
    ledger_nums = set(ANY_NUMBER.findall(ledger_text))
    claim_nums = {cid: set(ANY_NUMBER.findall(norm_num(c.get("text", "") + " " + (c.get("quote") or "")))) for cid, c in claims.items()}

    prev_layout = None
    layouts = []
    for i, s in enumerate(p.slides, 1):
        tag = f"slide {i}" + (f" #{s['id']}" if s["id"] else "")
        visible = ws(" ".join(s["text"]))
        notes = ws(" ".join(s["notes"]))
        everything = visible + " " + notes + " " + (s["title_attr"] or "")

        for m in DASHES.finditer(everything):
            ctx = everything[max(0, m.start() - 30): m.end() + 30]
            errors.append(f"{tag}: dash character '{m.group()}' in \"{ctx}\". Rewrite as two sentences, a comma, or a colon")
        for m in SPACED_HYPHEN.finditer(everything):
            ctx = everything[max(0, m.start() - 30): m.end() + 30]
            errors.append(f"{tag}: spaced hyphen used as a dash in \"{ctx}\"")
        if EMOJI.search(everything):
            errors.append(f"{tag}: emoji in slide text. Use type, rules, or drawn marks instead")
        for cid in s["claims"]:
            if brief and cid not in claims:
                errors.append(f"{tag}: data-claims references unknown claim '{cid}'")

        if brief:
            own = set().union(*(claim_nums.get(cid, set()) for cid in s["claims"])) if s["claims"] else None
            for n in NUMBER.findall(visible + " " + notes):
                k = norm_num(n)
                if k not in ledger_nums:
                    errors.append(f"{tag}: number '{n}' is not in any claim in brief.json. Add a sourced claim, or mark decorative numerals (folios, step markers) with data-deco")
                elif own is not None and k not in own:
                    warnings.append(f"{tag}: number '{n}' is in the ledger but not in this slide's data-claims. Cite the claim it comes from")

        if not s["layout"]:
            errors.append(f"{tag}: missing data-layout (needed for the variety check)")
        elif s["layout"] == prev_layout:
            warnings.append(f"{tag}: same data-layout '{s['layout']}' as the previous slide")
        prev_layout = s["layout"]
        layouts.append(s["layout"])

        title = ws(s["headings"][0] if s["headings"] else (s["title_attr"] or ""))
        if title.lower().strip(" .!:") in GENERIC_TITLES:
            warnings.append(f"{tag}: generic title '{title}'. Say the point instead")
        low = (visible + " " + notes).lower()
        for ph in BANNED_PHRASES:
            if re.search(r"\b" + re.escape(ph) + r"\b", low):
                warnings.append(f"{tag}: stock phrase '{ph}'")
        if NOT_X_BUT_Y.search(visible + " " + notes):
            warnings.append(f"{tag}: 'not X but Y' construction. State the point directly")
        if "!" in visible:
            warnings.append(f"{tag}: exclamation mark on a slide")
        raw = "\n".join(s["text"] + s["notes"])
        for sent in (x for chunk in raw.split("\n") for x in re.split(r"(?<=[.?!])\s+", ws(chunk))):
            if len(sent.split()) > 24:
                warnings.append(f"{tag}: long sentence ({len(sent.split())} words): \"{sent[:70]}...\"")
        if not s["has_notes"]:
            warnings.append(f"{tag}: no <aside class=\"notes\"> speaker notes")
        if s["anims"] > 9:
            warnings.append(f"{tag}: {s['anims']} animated elements. Choreograph fewer, larger moves")

    n = len(p.slides)
    if n >= 6 and len(set(layouts)) < min(4, n // 3 + 1):
        warnings.append(f"deck: only {len(set(layouts))} distinct layouts across {n} slides")
    return p, errors, warnings, verify


def gsap_tags(offline):
    if not offline:
        return "\n".join(f'<script src="{GSAP_CDN}{f}"></script>' for f in GSAP_FILES)
    vendor = SKILL / "assets" / "vendor"
    vendor.mkdir(exist_ok=True)
    out = []
    for f in GSAP_FILES:
        fp = vendor / f"{GSAP_VER}-{f}"
        if not fp.exists():
            with urllib.request.urlopen(GSAP_CDN + f, timeout=30) as r:
                fp.write_bytes(r.read())
        out.append(f"<script>{fp.read_text(encoding='utf-8')}</script>")
    return "\n".join(out)


def assemble(src, brief, offline):
    css = (SKILL / "assets" / "engine.css").read_text(encoding="utf-8")
    js = (SKILL / "assets" / "engine.js").read_text(encoding="utf-8")
    payload = {}
    if brief:
        payload = {
            "sources": {s["id"]: {"ref": s.get("ref"), "kind": s.get("kind")} for s in brief.get("sources", [])},
            "claims": {c["id"]: {"text": c.get("text"), "source": c.get("source"), "locator": c.get("locator")} for c in brief.get("claims", [])},
        }
    data = json.dumps(payload).replace("<", "\\u003c")
    head = f'<style id="kx-engine">\n{css}\n</style>'
    m = re.search(r"<head[^>]*>", src, re.I)
    if not m:
        raise SystemExit("build: no <head> in source")
    meta = re.search(r"<meta[^>]+charset[^>]*>", src[m.end():], re.I)
    at = m.end() + (meta.end() if meta and meta.start() < 200 else 0)
    src = src[:at] + "\n" + head + src[at:]
    tail = f'<script type="application/json" id="kx-sources">{data}</script>\n{gsap_tags(offline)}\n<script>\n{js}\n</script>\n'
    i = src.lower().rfind("</body>")
    if i < 0:
        raise SystemExit("build: no </body> in source")
    return src[:i] + tail + src[i:]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("-o", "--out")
    ap.add_argument("--brief")
    ap.add_argument("--offline", action="store_true", help="inline GSAP so the deck runs without network")
    ap.add_argument("--lint-only", action="store_true")
    ap.add_argument("--force", action="store_true", help="write even with errors (never for delivery)")
    a = ap.parse_args()
    for stream in (sys.stdout, sys.stderr):
        try: stream.reconfigure(encoding="utf-8")
        except Exception: pass

    src_path = Path(a.src).resolve()
    src = src_path.read_text(encoding="utf-8")
    brief_path = Path(a.brief).resolve() if a.brief else src_path.parent / "brief.json"
    brief = json.loads(brief_path.read_text(encoding="utf-8")) if brief_path.exists() else None

    p, errors, warnings, verify = lint(src, brief, brief_path.parent)
    print(f"stagecraft build: {len(p.slides)} slides, {len(errors)} errors, {len(warnings)} warnings")
    for e in errors: print("  ERROR  " + e)
    for w in warnings: print("  warn   " + w)
    if verify:
        print("  Claims from model knowledge (tell the user to verify):")
        for v in verify: print("    - " + v)
    if a.lint_only:
        sys.exit(1 if errors else 0)
    if errors and not a.force:
        print("Not written. Fix the errors and rebuild.")
        sys.exit(1)
    out = Path(a.out).resolve() if a.out else src_path.with_name(src_path.name.replace(".src.html", ".html") if ".src." in src_path.name else src_path.stem + ".built.html")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(assemble(src, brief, a.offline), encoding="utf-8")
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
