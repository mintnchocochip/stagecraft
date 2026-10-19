#!/usr/bin/env python3
"""Small helpers for native, editable PPTX decks with python-pptx (pip install python-pptx).

Coordinates are pixels on the same 1920x1080 canvas the HTML decks use. Import it from your build script:

    import sys; sys.path.insert(0, r"<skill>/scripts")
    from pptx_kit import Deck
    d = Deck(theme)                      # theme: bg, fg, muted, accent, display, body, mono (installed font names)
    s = d.slide(transition="morph")      # none | fade | push | morph ; pick from the motion level (see references/pptx.md)
    d.text(s, 96, 120, 1100, 300, "Title", font="display", size=120)
    d.save("deck.pptx")                  # refuses to save if any text has dash characters or emoji
"""
import re
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.util import Emu, Pt

PX = 6350  # EMU per px on a 1920-wide 13.333in slide
DASHES = re.compile(r"[‒–—―]|\w\s+-{1,2}\s+\w")
EMOJI = re.compile("[\U0001F300-\U0001FAFF\U00002600-\U000027BF\U0001F000-\U0001F2FF️]")

TRANSITIONS = {
    "fade": '<p:transition %s spd="med"><p:fade/></p:transition>' % nsdecls("p"),
    "push": '<p:transition %s spd="slow"><p:push dir="u"/></p:transition>' % nsdecls("p"),
    "morph": (
        '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" %s>'
        '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
        '<p:transition xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" spd="slow" p14:dur="1250">'
        '<p159:morph option="byObject"/></p:transition></mc:Choice>'
        '<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>' % nsdecls("p")
    ),
}


def px(v):
    return Emu(int(round(v * PX)))


def rgb(hexstr):
    return RGBColor.from_string(hexstr.lstrip("#").upper())


class Deck:
    def __init__(self, theme):
        self.t = theme
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = px(1920), px(1080)
        self.blank = self.prs.slide_layouts[6]

    def font(self, key):
        return self.t.get(key, key)

    def slide(self, bg=None, transition="none"):
        s = self.prs.slides.add_slide(self.blank)
        fill = s.background.fill
        fill.solid(); fill.fore_color.rgb = rgb(bg or self.t["bg"])
        if transition in TRANSITIONS:
            s._element.insert(2, parse_xml(TRANSITIONS[transition]))
        return s

    def text(self, s, x, y, w, h, txt, font="body", size=28, color=None, bold=False, italic=False,
             align="left", valign="top", line=1.15, tracking=0, name=None, upper=False):
        """txt may be a string or a list of strings (one paragraph each). Returns the shape."""
        tb = s.shapes.add_textbox(px(x), px(y), px(w), px(h))
        if name: tb.name = name
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[valign]
        for i, para in enumerate([txt] if isinstance(txt, str) else txt):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
            p.line_spacing = line
            r = p.add_run(); r.text = para.upper() if upper else para
            f = r.font
            f.name = self.font(font); f.size = Pt(size * 0.75); f.bold = bold; f.italic = italic
            f.color.rgb = rgb(color or self.t["fg"])
            if tracking:
                r._r.get_or_add_rPr().set("spc", str(int(tracking * 100)))
        return tb

    def rect(self, s, x, y, w, h, fill=None, line=None, line_w=2, name=None, shape=MSO_SHAPE.RECTANGLE):
        r = s.shapes.add_shape(shape, px(x), px(y), px(w), px(h))
        if name: r.name = name
        if fill: r.fill.solid(); r.fill.fore_color.rgb = rgb(fill)
        else: r.fill.background()
        if line: r.line.color.rgb = rgb(line); r.line.width = px(line_w)
        else: r.line.fill.background()
        r.shadow.inherit = False
        return r

    def rule(self, s, x, y, w, color=None, weight=2):
        return self.rect(s, x, y, w, weight, fill=color or self.t["fg"])

    def node(self, s, x, y, w, h, label, sub=None, fill=None, fg=None, line=None, name=None):
        box = self.rect(s, x, y, w, h, fill=fill or self.t["bg"], line=line or self.t["fg"], name=name)
        size = max(14, min(24, int((w - 48) / (0.62 * max(1, len(label))))))  # mono advance is about 0.6em
        self.text(s, x + 24, y + 22, w - 48, 40, label, font="mono", size=size, color=fg)
        if sub: self.text(s, x + 24, y + 62, w - 48, 40, sub, font="body", size=19, color=self.t.get("muted"))
        return box

    def edge(self, s, a, b, color=None, weight=2.5, curved=True, a_site=3, b_site=1):
        """Connect shape a to shape b with an arrow. Sites: 0 top, 1 left, 2 bottom, 3 right."""
        c = s.shapes.add_connector(MSO_CONNECTOR.CURVE if curved else MSO_CONNECTOR.STRAIGHT, 0, 0, 0, 0)
        c.begin_connect(a, a_site); c.end_connect(b, b_site)
        c.line.color.rgb = rgb(color or self.t["fg"]); c.line.width = px(weight)
        c.line._get_or_add_ln().append(parse_xml('<a:tailEnd %s type="triangle" w="med" len="med"/>' % nsdecls("a")))
        return c

    def chart(self, s, kind, categories, series, x, y, w, h, colors=None, labels=True, number_format='0.##', size=20):
        """kind: bar (horizontal) | column | line | donut. series: {"name": [values...]}. Native and editable."""
        ct = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
              "line": XL_CHART_TYPE.LINE_MARKERS, "donut": XL_CHART_TYPE.DOUGHNUT}[kind]
        cd = CategoryChartData(); cd.categories = categories
        for name, vals in series.items(): cd.add_series(name, vals)
        ch = s.shapes.add_chart(ct, px(x), px(y), px(w), px(h), cd).chart
        colors = colors or [self.t["fg"], self.t["accent"], self.t.get("muted", "888888")]
        ch.font.name = self.font("mono"); ch.font.size = Pt(size * 0.75); ch.font.color.rgb = rgb(self.t["fg"])
        ch.has_legend = len(series) > 1 or kind == "donut"
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
        plot = ch.plots[0]
        if kind == "donut":
            for i, pt in enumerate(plot.series[0].points):
                pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(colors[i % len(colors)])
        else:
            if hasattr(plot, "gap_width"): plot.gap_width = 60
            for i, ser in enumerate(plot.series):
                if kind == "line":
                    ser.format.line.color.rgb = rgb(colors[i % len(colors)]); ser.format.line.width = px(4); ser.smooth = False
                else:
                    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(colors[i % len(colors)])
            va, ca = ch.value_axis, ch.category_axis
            va.has_major_gridlines = False; va.visible = not labels
            ca.format.line.color.rgb = rgb(self.t["fg"]); ca.tick_labels.font.size = Pt(size * 0.75)
        if labels:
            plot.has_data_labels = True
            dl = plot.data_labels; dl.number_format = number_format; dl.number_format_is_linked = False
            dl.font.size = Pt(size * 0.75)
            if kind in ("bar", "column"): dl.position = XL_LABEL_POSITION.OUTSIDE_END
        return ch

    def image(self, s, path, x, y, w=None, h=None, name=None):
        pic = s.shapes.add_picture(str(path), px(x), px(y), px(w) if w else None, px(h) if h else None)
        if name: pic.name = name
        return pic

    def notes(self, s, txt):
        s.notes_slide.notes_text_frame.text = txt

    def check(self):
        bad = []
        for i, s in enumerate(self.prs.slides, 1):
            texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
            if s.has_notes_slide: texts.append(s.notes_slide.notes_text_frame.text)
            for t in texts:
                if DASHES.search(t): bad.append(f"slide {i}: dash in \"{t[:60]}\"")
                if EMOJI.search(t): bad.append(f"slide {i}: emoji in \"{t[:60]}\"")
        return bad

    def save(self, path):
        bad = self.check()
        if bad:
            raise SystemExit("pptx not saved:\n  " + "\n  ".join(bad))
        self.prs.save(path)
        print(f"wrote {path} ({len(self.prs.slides)} slides)")
